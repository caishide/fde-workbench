"""WP04 文档可信交付回归钉（D3/D4/D5/D9/D10/E1）。

- D3/E1：AI 生成字段在成文中带「AI 生成，请核实」标注（S10 纪律），人工覆盖后不标
- D10：同一冻结版本重复生成不实时读库，AI 字段与 generated_date 均来自首代快照
- D4：draft 版本不可签发对外下载 URL（含普通用户/跨项目访问路径）
- D5：下载签发落 document_version_delivered 审计（操作人/版本/sha256/时间）
- D9：文档模板 published 版本拒绝再发布/修改/删除，草稿仍可试渲染
"""
import io

from docx import Document
from flask import current_app
from sqlalchemy import select

from fde_api.documents.generator import handle_document_generation
from fde_api.workbench.models import OperationEvent
from fde_api.documents.qa import extract_docx_xml_text
from fde_api.extensions import object_storage
from fde_api.files.models import ProjectFileVersion
from test_document_generator import _seed_generation
from test_document_versions_api import (
    _docx_bytes,
    _manual_upload,
    _seed_document,
    authorized_client_factory,
    engineer_client,
    leader_client,
)


class _FakeDocumentAI:
    """complete_json 形状的最小替身；payload 可变以验证冻结语义。"""

    payload = {"fields": {"current_situation": "排产依赖手工台账"}}

    def complete_json(self, *, messages, request_id, max_tokens):
        return dict(self.payload)


def _build_docx(builder) -> bytes:
    doc = Document()
    builder(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _ai_placeholder_builder(doc) -> None:
    doc.add_paragraph("现状：【待填写：现状描述】")
    doc.add_paragraph("生成日期：【待填写：生成日期】")


def _ai_placeholder_template_bytes() -> bytes:
    return _build_docx(_ai_placeholder_builder)


def _stored_version_text(db_session, version) -> str:
    db_session.refresh(version)
    file_version = db_session.scalar(
        select(ProjectFileVersion).where(
            ProjectFileVersion.id == version.docx_file_version_id
        )
    )
    assert file_version is not None
    return extract_docx_xml_text(object_storage.current.open_stream(file_version.storage_key))


def _mark_ungenerated(db_session, version) -> None:
    """模拟「重新生成」触发：清空 sha256 即绕过幂等短路，允许再次完整渲染。

    docx_file_version_id 是 NOT NULL 外键，不能置空。
    """
    version.sha256 = ""
    db_session.commit()


# ---------------------------------------------------------------------------
# D3/E1：AI 标注
# ---------------------------------------------------------------------------

def test_ai_generated_field_is_annotated_in_rendered_docx(db_session):
    seeded = _seed_generation(db_session, template_bytes=_ai_placeholder_template_bytes())
    seeded["document"].opportunity_scope_json = {"name": "排产优化"}
    db_session.commit()
    version = seeded["version"]
    current_app.extensions["fde_api_document_generation_deepseek"] = _FakeDocumentAI()
    try:
        result = handle_document_generation(version_id=version.id)
        assert result["status"] == "succeeded", result
        text = _stored_version_text(db_session, version)
        assert "排产依赖手工台账（AI 生成，请核实）" in text
    finally:
        current_app.extensions.pop("fde_api_document_generation_deepseek", None)


def test_manual_override_replaces_ai_value_without_annotation(db_session):
    seeded = _seed_generation(db_session, template_bytes=_ai_placeholder_template_bytes())
    version = seeded["version"]
    version.content_snapshot_json = {
        "field_overrides": {"current_situation": "人工核对后的现状"},
    }
    db_session.commit()
    current_app.extensions["fde_api_document_generation_deepseek"] = _FakeDocumentAI()
    try:
        result = handle_document_generation(version_id=version.id)
        assert result["status"] == "succeeded", result
        text = _stored_version_text(db_session, version)
        assert "人工核对后的现状" in text
        assert "排产依赖手工台账" not in text
        assert "（AI 生成，请核实）" not in text
    finally:
        current_app.extensions.pop("fde_api_document_generation_deepseek", None)


# ---------------------------------------------------------------------------
# D10：冻结可重现
# ---------------------------------------------------------------------------

def test_regeneration_reuses_frozen_ai_fields_not_live_ai(db_session):
    seeded = _seed_generation(db_session, template_bytes=_ai_placeholder_template_bytes())
    seeded["document"].opportunity_scope_json = {"name": "排产优化"}
    db_session.commit()
    version = seeded["version"]
    fake = _FakeDocumentAI()
    current_app.extensions["fde_api_document_generation_deepseek"] = fake
    try:
        first = handle_document_generation(version_id=version.id)
        assert first["status"] == "succeeded", first
        first_text = _stored_version_text(db_session, version)
        first_date_line = next(
            line for line in first_text.splitlines() if "生成日期" in line
        )

        # 首代之后：AI 输出整体变化（换模型/换上下文），已冻结版本不得漂移
        fake.payload = {"fields": {"current_situation": "生成后新写的现状"}}
        _mark_ungenerated(db_session, version)

        second = handle_document_generation(version_id=version.id)
        assert second["status"] == "succeeded", second
        second_text = _stored_version_text(db_session, version)

        assert "排产依赖手工台账（AI 生成，请核实）" in second_text
        assert "生成后新写的现状" not in second_text
        second_date_line = next(
            line for line in second_text.splitlines() if "生成日期" in line
        )
        assert second_date_line == first_date_line
    finally:
        current_app.extensions.pop("fde_api_document_generation_deepseek", None)


# ---------------------------------------------------------------------------
# D4/D5：下载门禁与交付审计
# ---------------------------------------------------------------------------

def test_draft_download_gate_and_delivery_audit(db_session, leader_client, engineer_client):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    project_id = seeded["project"].id
    document = seeded["document"]
    base = f"/api/v1/projects/{project_id}/documents"

    upload = _manual_upload(
        engineer_client, project_id, document.id, document.version,
        _docx_bytes(text="交付版本"),
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)
    v2 = upload.json["data"]

    # D4：普通项目成员对 draft 请求下载 → 409（不是 200）
    draft = engineer_client.get(f"{base}/{v2['id']}/download-url")
    assert draft.status_code == 409
    assert draft.json["error"]["code"] == "document_version_not_confirmed"

    assert leader_client.post(
        f"{base}/{v2['id']}/confirm", json={"version": v2["document_version"]}
    ).status_code == 200

    # D4：项目成员（非管理员）对 confirmed 版本可下载
    delivered = engineer_client.get(f"{base}/{v2['id']}/download-url")
    assert delivered.status_code == 200, delivered.get_data(as_text=True)

    # D5：签发落审计——操作人/版本/sha256/时间可查
    event = db_session.scalar(
        select(OperationEvent)
        .where(
            OperationEvent.event_type == "document_version_delivered",
            OperationEvent.project_id == project_id,
        )
        .order_by(OperationEvent.occurred_at.desc())
    )
    assert event is not None
    assert event.actor_user_id == engineer.id
    assert event.changes["version_id"] == v2["id"]
    assert event.changes["sha256"]
    assert event.occurred_at is not None


def test_download_rejected_for_non_member_and_cross_project(
    db_session, leader_client, engineer_client, authorized_client_factory
):
    leader = leader_client.user
    engineer = engineer_client.user
    outsider = authorized_client_factory("fde_engineer")
    seeded = _seed_document(db_session, leader, engineer)
    project_id = seeded["project"].id
    document = seeded["document"]
    base = f"/api/v1/projects/{project_id}/documents"

    upload = _manual_upload(
        engineer_client, project_id, document.id, document.version,
        _docx_bytes(text="确认后交付"),
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)
    v2 = upload.json["data"]
    assert leader_client.post(
        f"{base}/{v2['id']}/confirm", json={"version": v2["document_version"]}
    ).status_code == 200

    # 非项目成员：项目不可见 → 拒绝
    outsider_view = outsider.get(f"{base}/{v2['id']}/download-url")
    assert outsider_view.status_code == 403

    # 跨项目路径：文档不属于该路径上的项目 → 404
    cross = engineer_client.get(
        f"/api/v1/projects/00000000-0000-0000-0000-000000000000/documents/{v2['id']}/download-url"
    )
    assert cross.status_code == 404


# ---------------------------------------------------------------------------
# D9：文档模板发布锁
# ---------------------------------------------------------------------------

def test_published_document_template_rejects_republish_and_mutation(
    db_session, authorized_client_factory
):
    admin = authorized_client_factory("admin")
    headers = admin.headers
    base = "/api/v1/document-templates"
    upload = admin.client.post(
        f"{base}/sow/versions",
        data={"file": (io.BytesIO(_docx_bytes(text="sow 模板源")), "sow.docx")},
        headers=headers,
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)
    draft = upload.json["data"]

    # 草稿仍可编辑路径：试渲染可用
    test_render = admin.client.post(
        f"{base}/{draft['id']}/test", headers=headers
    )
    assert test_render.status_code == 200, test_render.get_data(as_text=True)

    publish = admin.client.post(
        f"{base}/{draft['id']}/publish", headers=headers
    )
    assert publish.status_code == 200, publish.get_data(as_text=True)

    # 已发布版本：再发布被拒（明确冲突错误）
    republish = admin.client.post(
        f"{base}/{draft['id']}/publish", headers=headers
    )
    assert republish.status_code == 409
    assert republish.json["error"]["code"] == "invalid_version_state"

    # 已发布版本：修改/删除没有服务端入口（PATCH/DELETE 直接 404/405）
    patch = admin.client.patch(
        f"{base}/{draft['id']}", json={}, headers=headers
    )
    assert patch.status_code in (404, 405)
    delete = admin.client.delete(
        f"{base}/{draft['id']}", headers=headers
    )
    assert delete.status_code in (404, 405)


def test_draft_document_template_can_be_deactivated_but_published_cannot(
    db_session, authorized_client_factory
):
    """D9 正反向：草稿可停用；已发布版本停用被拒（锁版硬拦截）。"""
    admin = authorized_client_factory("admin")
    headers = admin.headers
    base = "/api/v1/document-templates"

    draft = admin.client.post(
        f"{base}/sow/versions",
        data={"file": (io.BytesIO(_docx_bytes(text="可停用草稿")), "sow.docx")},
        headers=headers,
    ).json["data"]
    deactivated = admin.client.post(
        f"{base}/{draft['id']}/deactivate", headers=headers
    )
    assert deactivated.status_code == 200, deactivated.get_data(as_text=True)
    assert deactivated.json["data"]["status"] == "inactive"

    published = admin.client.post(
        f"{base}/sow/versions",
        data={"file": (io.BytesIO(_docx_bytes(text="锁版模板")), "sow.docx")},
        headers=headers,
    ).json["data"]
    assert admin.client.post(
        f"{base}/{published['id']}/publish", headers=headers
    ).status_code == 200

    locked = admin.client.post(f"{base}/{published['id']}/deactivate", headers=headers)
    assert locked.status_code == 409, locked.get_data(as_text=True)
    assert locked.json["error"]["code"] == "template_version_immutable"
