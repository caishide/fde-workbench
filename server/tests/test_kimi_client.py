import io
import json

import pytest
import requests

from fde_api.ai.kimi import KimiError, KimiFileClient


class FakeResponse:
    def __init__(self, status_code=200, payload=None, text="", content=None, stream_lines=None):
        self.status_code = status_code
        self._payload = payload
        self.text = text
        self.content = text.encode("utf-8") if content is None else content
        self._stream_lines = stream_lines or []

    def json(self):
        if isinstance(self._payload, Exception):
            raise self._payload
        return self._payload

    def iter_lines(self, decode_unicode=False):
        return iter(self._stream_lines)


class FakeSession:
    def __init__(self, *, upload=None, content=None, completion=None, error=None):
        self.upload = upload
        self.content = content
        self.completion = completion
        self.error = error
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append(("post", url, kwargs))
        if self.error:
            raise self.error
        return self.upload if url.endswith("/files") else self.completion

    def get(self, url, **kwargs):
        self.calls.append(("get", url, kwargs))
        return self.content

    def delete(self, url, **kwargs):
        self.calls.append(("delete", url, kwargs))
        return FakeResponse(payload={"deleted": True})


def make_client(session, api_key="secret"):
    return KimiFileClient(base_url="https://api.moonshot.example/v1", api_key=api_key, model="kimi-k3", timeout_seconds=30, session=session)


def test_uploads_docx_uses_kimi_content_and_cleans_up_file():
    result = {"customer_vision": "长期目标"}
    session = FakeSession(
        upload=FakeResponse(payload={"id": "file-1", "status": "ready"}),
        content=FakeResponse(text="Kimi 原生提取的文档内容"),
        completion=FakeResponse(stream_lines=[
            ("data: " + json.dumps({"choices": [{"delta": {"content": json.dumps(result, ensure_ascii=False)}}]}, ensure_ascii=False)).encode("utf-8"),
            b"data: [DONE]",
        ]),
    )
    messages_seen = []

    parsed = make_client(session).analyze_docx_json(
        stream=io.BytesIO(b"docx"), filename="预调研表.docx", request_id="job-1",
        build_messages=lambda content: messages_seen.append(content) or [{"role": "system", "content": content}],
    )

    assert parsed == result
    assert messages_seen == ["Kimi 原生提取的文档内容"]
    upload = session.calls[0]
    assert upload[2]["data"] == {"purpose": "file-extract"}
    assert upload[2]["files"]["file"][0] == "预调研表.docx"
    chat = next(call for call in session.calls if call[1].endswith("/chat/completions"))
    assert chat[2]["json"]["model"] == "kimi-k3"
    assert chat[2]["json"]["response_format"] == {"type": "json_object"}
    assert chat[2]["json"]["reasoning_effort"] == "low"
    assert chat[2]["json"]["max_completion_tokens"] == 32768
    assert chat[2]["json"]["stream"] is True
    assert chat[2]["stream"] is True
    assert "temperature" not in chat[2]["json"]
    assert session.calls[-1][0:2] == ("delete", "https://api.moonshot.example/v1/files/file-1")


def test_missing_key_stops_before_upload():
    session = FakeSession()
    with pytest.raises(KimiError) as caught:
        make_client(session, api_key=None).analyze_docx_json(stream=io.BytesIO(), filename="a.docx", request_id="job", build_messages=lambda _: [])
    assert caught.value.code == "ai_service_not_configured"
    assert session.calls == []


def test_timeout_has_chinese_public_error():
    session = FakeSession(error=requests.Timeout("private timeout"))
    with pytest.raises(KimiError) as caught:
        make_client(session).analyze_docx_json(stream=io.BytesIO(), filename="a.docx", request_id="job", build_messages=lambda _: [])
    assert caught.value.code == "ai_timeout"
    assert caught.value.public_message == "AI 分析超时，请稍后重试。"


def test_invalid_completion_is_reported_and_remote_file_is_removed():
    session = FakeSession(
        upload=FakeResponse(payload={"id": "file-2", "status": "ready"}),
        content=FakeResponse(text="content"),
        completion=FakeResponse(stream_lines=["data: [DONE]"]),
    )
    with pytest.raises(KimiError) as caught:
        make_client(session).analyze_docx_json(stream=io.BytesIO(), filename="a.docx", request_id="job", build_messages=lambda _: [])
    assert caught.value.code == "ai_empty_response"
    assert session.calls[-1][0] == "delete"


def _docx_bytes(paragraph: str = "FDE 预调研：客户希望 AI 落地", table_cell: str = "50万") -> io.BytesIO:
    from docx import Document

    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "预算"
    table.rows[0].cells[1].text = table_cell
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_openai_compatible_client(session):
    """MiniMax 等通用 OpenAI 兼容端点：本地提取 DOCX + max_tokens。"""
    return KimiFileClient(
        base_url="https://api.minimaxi.example/v1",
        api_key="secret",
        model="MiniMax-M3",
        timeout_seconds=30,
        session=session,
        local_extract=True,
        openai_compatible=True,
    )


def test_local_extract_skips_upload_and_sends_docx_text():
    """MiniMax 没有 file-extract 接口：不得调用 /files，改为本地提取正文。"""
    result = {"customer_vision": "长期目标"}
    session = FakeSession(
        completion=FakeResponse(stream_lines=[
            ("data: " + json.dumps({"choices": [{"delta": {"content": json.dumps(result, ensure_ascii=False)}}]}, ensure_ascii=False)).encode("utf-8"),
            b"data: [DONE]",
        ]),
    )
    messages_seen = []

    parsed = make_openai_compatible_client(session).analyze_docx_json(
        stream=_docx_bytes(), filename="预调研表.docx", request_id="job-1",
        build_messages=lambda content: messages_seen.append(content) or [{"role": "system", "content": content}],
    )

    assert parsed == result
    # 段落与表格内容都被提取出来
    assert len(messages_seen) == 1
    assert "FDE 预调研" in messages_seen[0]
    assert "50万" in messages_seen[0]
    # 全程没有 /files 上传与删除
    assert not any(call[1].endswith("/files") for call in session.calls)
    assert not any(call[0] == "delete" for call in session.calls)


def test_openai_compatible_payload_uses_max_tokens_and_no_reasoning_effort():
    """MiniMax 侧不识别 reasoning_effort / max_completion_tokens。"""
    session = FakeSession(
        completion=FakeResponse(stream_lines=[
            ("data: " + json.dumps({"choices": [{"delta": {"content": "{}"}}]}, ensure_ascii=False)).encode("utf-8"),
            b"data: [DONE]",
        ]),
    )

    make_openai_compatible_client(session).analyze_docx_json(
        stream=_docx_bytes(), filename="a.docx", request_id="job-1",
        build_messages=lambda content: [{"role": "user", "content": content}],
    )

    chat = next(call for call in session.calls if call[1].endswith("/chat/completions"))
    payload = chat[2]["json"]
    assert chat[1] == "https://api.minimaxi.example/v1/chat/completions"
    assert payload["model"] == "MiniMax-M3"
    assert payload["max_tokens"] == 32768
    assert "max_completion_tokens" not in payload
    assert "reasoning_effort" not in payload
    assert payload["response_format"] == {"type": "json_object"}


def test_local_extract_reports_unparsable_docx():
    """DOCX 解析失败要给中文错误，不能抛裸异常。

    可重试性按失败性质区分：本地解析是字节的纯函数，同一份坏文件每次都
    解析失败——重试无意义，故不可重试（评审 G5-① F3 的结论）。
    """
    session = FakeSession()
    with pytest.raises(KimiError) as caught:
        make_openai_compatible_client(session).analyze_docx_json(
            stream=io.BytesIO(b"not-a-docx"), filename="坏文件.docx", request_id="job-1",
            build_messages=lambda content: [{"role": "user", "content": content}],
        )
    assert caught.value.code == "kimi_file_extract_failed"
    assert caught.value.retryable is False
    assert not any(call[1].endswith("/chat/completions") for call in session.calls)
